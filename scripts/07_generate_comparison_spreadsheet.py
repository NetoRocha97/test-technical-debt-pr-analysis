
import json
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill,
    Font,
    Alignment,
    Border,
    Side
)
from openpyxl.utils import get_column_letter

# =====================================================
# FILE PATHS
# =====================================================

MANUAL_FILE = "data/manual_validation.xlsx"

LLM_FILE = "data/llm_results_validation.json"

OUTPUT_FILE = "results/validação/manual_validation_with_llm.xlsx"

# =====================================================
# LOAD DATA
# =====================================================

print("Loading manual spreadsheet...")

manual = pd.read_excel(MANUAL_FILE)

manual.columns = manual.columns.str.strip()

manual = manual.loc[:, manual.columns != ""]

manual = manual.dropna(subset=["repo", "pr_number"])

manual = manual[
    manual["repo"].astype(str).str.lower() != "repo"
]

manual = manual[
    manual["pr_number"].astype(str).str.lower() != "pr_number"
]

manual["pr_number"] = (
    manual["pr_number"]
    .astype(str)
    .str.strip()
    .str.replace("#", "", regex=False)
)

manual["pr_number"] = pd.to_numeric(
    manual["pr_number"],
    errors="coerce"
)

manual = manual.dropna(subset=["repo", "pr_number"])

manual["pr_number"] = manual["pr_number"].astype(int)

manual = manual.dropna(subset=["repo", "pr_number"])

print("Loading LLM results...")

with open(
    LLM_FILE,
    encoding="utf8"
) as f:

    llm = json.load(f)

# =====================================================
# INDEX LLM RESULTS
# =====================================================

llm_dict = {}

for pr in llm:

    key = (
        pr["repo"],
        int(pr["pr_number"])
    )

    llm_dict[key] = {

        "pr_objective":
            pr["primary_pr_objective"]["value"],

        "suggestion_type":
            pr["suggestion_type"]["value"],

        "agreement":
            pr["agreement"]["value"],

        "resolution":
            pr["resolution"]["value"],

        "test_td_mentioned":
            pr["test_td_mentioned"]["value"],

        "mention_type":
            pr["mention_type"]["value"],

        "test_td_type":
            pr["test_td_type"]["value"]

    }

print(f"Indexed results: {len(llm_dict)}")

# =====================================================
# CREATE EXCEL WORKBOOK
# =====================================================

wb = Workbook()

ws = wb.active

ws.title = "Validation"

# =====================================================
# STYLES
# =====================================================

BLUE = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

GREEN = PatternFill(
    "solid",
    fgColor="D9EAD3"
)

YELLOW = PatternFill(
    "solid",
    fgColor="FFF2CC"
)

ORANGE = PatternFill(
    "solid",
    fgColor="F9CB9C"
)

BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

BOLD = Font(
    bold=True
)

CENTER = Alignment(
    horizontal="center",
    vertical="center"
)

# =====================================================
# MAIN HEADER
# =====================================================

groups = [

    ("A1:B1", "Identification", BLUE),

    ("C1:E1", "RQ1 - Primary PR Objective", GREEN),

    ("F1:H1", "RQ2 - Suggestion Type", YELLOW),

    ("I1:K1", "RQ2 - Agreement", YELLOW),

    ("L1:N1", "RQ2 - Resolution", YELLOW),

    ("O1:Q1", "RQ4 - Test TD Mention", ORANGE),

    ("R1:T1", "RQ4 - Mention Type", ORANGE),

    ("U1:W1", "RQ4 - Test TD Type", ORANGE)

]

for cell_range, title, color in groups:

    ws.merge_cells(cell_range)

    cell = ws[
        cell_range.split(":")[0]
    ]

    cell.value = title

    cell.fill = color

    cell.font = BOLD

    cell.alignment = CENTER

    cell.border = BORDER

# =====================================================
# SECOND HEADER ROW
# =====================================================

headers = [

    "Repository",
    "PR",

    "Manual",
    "LLM",
    "Match",

    "Manual",
    "LLM",
    "Match",

    "Manual",
    "LLM",
    "Match",

    "Manual",
    "LLM",
    "Match",

    "Manual",
    "LLM",
    "Match",

    "Manual",
    "LLM",
    "Match",

    "Manual",
    "LLM",
    "Match"

]

header_colors = [

    BLUE,
    BLUE,

    GREEN,
    GREEN,
    GREEN,

    YELLOW,
    YELLOW,
    YELLOW,

    YELLOW,
    YELLOW,
    YELLOW,

    YELLOW,
    YELLOW,
    YELLOW,

    ORANGE,
    ORANGE,
    ORANGE,

    ORANGE,
    ORANGE,
    ORANGE,

    ORANGE,
    ORANGE,
    ORANGE

]

for column, title in enumerate(headers, start=1):

    cell = ws.cell(row=2, column=column)

    cell.value = title

    cell.font = BOLD

    cell.alignment = CENTER

    cell.fill = header_colors[column - 1]

    cell.border = BORDER

# =====================================================
# DATA ROWS
# =====================================================

row_number = 3

for _, row in manual.iterrows():

    repo = row["repo"]

    pr_number = row["pr_number"]

    key = (
        repo,
        pr_number
    )

    llm_data = llm_dict.get(key, {})

    values = [

        repo,
        pr_number,

        row["objetivo_pr"],
        llm_data.get("pr_objective", ""),
        "",

        row["tipo_sugestao"],
        llm_data.get("suggestion_type", ""),
        "",

        row["acordo"],
        llm_data.get("agreement", ""),
        "",

        row["resolucao"],
        llm_data.get("resolution", ""),
        "",

        row["dt_teste_mencionada"],
        llm_data.get("test_td_mentioned", ""),
        "",

        row["forma_dt_menc"],
        llm_data.get("mention_type", ""),
        "",

        row["tipo_dt_teste"],
        llm_data.get("test_td_type", ""),
        ""

    ]

    for column, value in enumerate(values, start=1):

        cell = ws.cell(
            row=row_number,
            column=column
        )

        cell.value = value

        cell.alignment = CENTER

        cell.border = BORDER

    row_number += 1

# =====================================================
# AUTOMATIC COLUMN WIDTH ADJUSTMENT
# =====================================================

for column_cells in ws.columns:

    column_letter = get_column_letter(
        column_cells[0].column
    )

    max_length = 0

    for cell in column_cells:

        if cell.value is not None:

            max_length = max(
                max_length,
                len(str(cell.value))
            )

    ws.column_dimensions[column_letter].width = min(
        max_length + 3,
        35
    )

for column in [
    "E",
    "H",
    "K",
    "N",
    "Q",
    "T",
    "W"
]:

    ws.column_dimensions[column].width = 10

# =====================================================
# WORKSHEET CONFIGURATION
# =====================================================

ws.freeze_panes = "A3"

ws.auto_filter.ref = f"A2:W{ws.max_row}"

# =====================================================
# SAVE WORKBOOK
# =====================================================

wb.save(OUTPUT_FILE)

print()

print("=" * 60)

print("SPREADSHEET GENERATED SUCCESSFULLY")

print("=" * 60)

print()

print(f"File saved to: {OUTPUT_FILE}")
